import json
import re
import sys
from collections import Counter

from docx import Document


def inspect(path: str) -> dict:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    non_empty = [text for text in paragraphs if text.strip()]
    keys = []
    headings = []
    for number, text in enumerate(paragraphs, 1):
        stripped = text.strip()
        match = re.match(r'^\s*"([^"\n]{1,60})"\s*:', text)
        if match:
            keys.append({"paragraph": number, "key": match.group(1), "preview": stripped[:220]})
        if stripped.startswith(("# ", "## ", "<Roleplay", "</Roleplay")):
            headings.append({"paragraph": number, "text": stripped[:220]})
    return {
        "path": path,
        "paragraphs": len(non_empty),
        "characters": len("\n".join(non_empty)),
        "tables": len(document.tables),
        "keys": keys,
        "headings": headings,
        "style_counts": Counter(
            paragraph.style.name if paragraph.style else "None"
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ),
    }


if __name__ == "__main__":
    print(json.dumps([inspect(path) for path in sys.argv[1:]], ensure_ascii=False, indent=2))
